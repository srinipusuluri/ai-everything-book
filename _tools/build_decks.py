#!/usr/bin/env python3
"""
build_decks.py — turn every slides/deck.json into a .pptx and a Marp-ready slides.md.

    ../.venv/bin/python _tools/build_decks.py            # build all modules
    ../.venv/bin/python _tools/build_decks.py 04-llm     # build one

Deck JSON schema
----------------
{
  "module": "01", "title": "...", "subtitle": "...", "accent": "RRGGBB",
  "slides": [
     {"type":"bullets", "title":"...", "bullets":[...], "notes":"..."},
     {"type":"table",   "title":"...", "columns":[...], "rows":[[...]], "notes":"..."},
     {"type":"quote",   "title":"...", "quote":"...", "bullets":[...], "notes":"..."},
     {"type":"section", "title":"...", "subtitle":"...", "notes":"..."}
  ]
}
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
W, H = Inches(13.333), Inches(7.5)          # 16:9
INK = RGBColor(0x1A, 0x1D, 0x29)
MUTED = RGBColor(0x6B, 0x72, 0x84)
PAPER = RGBColor(0xFF, 0xFF, 0xFF)
WASH = RGBColor(0xF6, 0xF7, 0xFB)
FONT = "Calibri"
MONO = "Consolas"


def hexc(s: str) -> RGBColor:
    return RGBColor.from_string(s.replace("#", "").upper())


def textbox(slide, l, t, w, h, *, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def para(tf, text, *, size, bold=False, color=INK, first=False,
         space_after=6, font=FONT, align=PP_ALIGN.LEFT, italic=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = font
    return p


def rect(slide, l, t, w, h, fill, line=None):
    from pptx.enum.shapes import MSO_SHAPE
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(0.75)
    sh.shadow.inherit = False
    return sh


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def set_notes(slide, text):
    if text:
        slide.notes_slide.notes_text_frame.text = text


def chrome(slide, accent, module, title, idx, total):
    """Common header/footer furniture for content slides."""
    rect(slide, 0, 0, W, Emu(int(Inches(0.09))), accent)
    tf = textbox(slide, Inches(0.7), Inches(0.52), Inches(11.9), Inches(0.9))
    para(tf, title, size=30, bold=True, first=True, space_after=0)
    rect(slide, Inches(0.7), Inches(1.42), Inches(1.0), Emu(int(Pt(3))), accent)
    ftf = textbox(slide, Inches(0.7), Inches(6.95), Inches(11.9), Inches(0.35))
    p = ftf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = f"Module {module}"
    r.font.size = Pt(10)
    r.font.color.rgb = MUTED
    r.font.name = FONT
    ntf = textbox(slide, Inches(11.6), Inches(6.95), Inches(1.0), Inches(0.35))
    p2 = ntf.paragraphs[0]
    p2.alignment = PP_ALIGN.RIGHT
    r2 = p2.add_run()
    r2.text = f"{idx} / {total}"
    r2.font.size = Pt(10)
    r2.font.color.rgb = MUTED
    r2.font.name = FONT


def slide_title(prs, deck, accent):
    s = blank(prs)
    rect(s, 0, 0, W, H, accent)
    rect(s, 0, Inches(5.55), W, Inches(1.95), PAPER)
    tf = textbox(s, Inches(0.9), Inches(2.0), Inches(11.5), Inches(2.6))
    para(tf, f"MODULE {deck['module']}", size=16, bold=True,
         color=RGBColor(0xFF, 0xFF, 0xFF), first=True, space_after=14)
    para(tf, deck["title"], size=54, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), space_after=10)
    para(tf, deck.get("subtitle", ""), size=20, color=RGBColor(0xE8, 0xEA, 0xF2))
    tf2 = textbox(s, Inches(0.9), Inches(6.05), Inches(11.5), Inches(1.0))
    para(tf2, "AI End-to-End Learning Track", size=18, bold=True, color=accent, first=True, space_after=4)
    para(tf2, "Notes · Code · Labs · Papers — see the module README", size=12, color=MUTED)
    set_notes(s, deck.get("notes", ""))
    return s


def slide_section(prs, sl, accent, module, idx, total):
    s = blank(prs)
    rect(s, 0, 0, W, H, WASH)
    rect(s, 0, 0, Inches(0.35), H, accent)
    tf = textbox(s, Inches(1.2), Inches(2.9), Inches(11.0), Inches(2.0))
    para(tf, sl["title"], size=44, bold=True, first=True, space_after=10)
    if sl.get("subtitle"):
        para(tf, sl["subtitle"], size=18, color=MUTED)
    set_notes(s, sl.get("notes", ""))


def slide_bullets(prs, sl, accent, module, idx, total):
    s = blank(prs)
    chrome(s, accent, module, sl["title"], idx, total)
    bullets = sl.get("bullets", [])
    size = 20 if len(bullets) <= 5 else 17
    tf = textbox(s, Inches(0.7), Inches(1.85), Inches(11.9), Inches(4.9))
    for i, b in enumerate(bullets):
        sub = b.startswith("- ")
        txt = b[2:] if sub else b
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(13 if not sub else 7)
        p.level = 1 if sub else 0
        r0 = p.add_run()
        r0.text = "—  " if sub else "▪  "
        r0.font.size = Pt(size - (3 if sub else 0))
        r0.font.color.rgb = accent
        r0.font.name = FONT
        r = p.add_run()
        r.text = txt
        r.font.size = Pt(size - (3 if sub else 0))
        r.font.color.rgb = INK if not sub else MUTED
        r.font.name = MONO if ("=" in txt and "<-" in txt) else FONT
    set_notes(s, sl.get("notes", ""))


def slide_table(prs, sl, accent, module, idx, total):
    s = blank(prs)
    chrome(s, accent, module, sl["title"], idx, total)
    cols, rows = sl["columns"], sl["rows"]
    nr, nc = len(rows) + 1, len(cols)
    top = Inches(1.9)
    height = min(Inches(4.8), Inches(0.42) * nr)
    gf = s.shapes.add_table(nr, nc, Inches(0.7), top, Inches(11.9), height).table
    widths = sl.get("widths")
    if widths and len(widths) == nc:
        total_w = Inches(11.9)
        for i, frac in enumerate(widths):
            gf.columns[i].width = Emu(int(total_w * frac))
    for j, c in enumerate(cols):
        cell = gf.cell(0, j)
        cell.text = ""
        cell.fill.solid()
        cell.fill.fore_color.rgb = accent
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        r = p.add_run()
        r.text = str(c)
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.name = FONT
    fs = 13 if len(rows) <= 6 else 11
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            cell = gf.cell(i, j)
            cell.text = ""
            cell.fill.solid()
            cell.fill.fore_color.rgb = PAPER if i % 2 else WASH
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            r = p.add_run()
            r.text = str(val)
            r.font.size = Pt(fs)
            r.font.color.rgb = INK
            r.font.bold = (j == 0)
            r.font.name = FONT
    set_notes(s, sl.get("notes", ""))


def slide_quote(prs, sl, accent, module, idx, total):
    s = blank(prs)
    chrome(s, accent, module, sl["title"], idx, total)
    rect(s, Inches(0.7), Inches(1.9), Inches(11.9), Inches(1.7), WASH)
    rect(s, Inches(0.7), Inches(1.9), Emu(int(Pt(5))), Inches(1.7), accent)
    tf = textbox(s, Inches(1.15), Inches(2.15), Inches(11.1), Inches(1.3), anchor=MSO_ANCHOR.MIDDLE)
    para(tf, sl["quote"], size=24, bold=True, italic=True, first=True, space_after=0)
    bl = sl.get("bullets", [])
    if bl:
        tf2 = textbox(s, Inches(0.7), Inches(3.95), Inches(11.9), Inches(2.6))
        for i, b in enumerate(bl):
            p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
            p.space_after = Pt(11)
            r0 = p.add_run()
            r0.text = "▪  "
            r0.font.size = Pt(17)
            r0.font.color.rgb = accent
            r0.font.name = FONT
            r = p.add_run()
            r.text = b
            r.font.size = Pt(17)
            r.font.color.rgb = INK
            r.font.name = FONT
    set_notes(s, sl.get("notes", ""))


def slide_end(prs, deck, accent):
    s = blank(prs)
    rect(s, 0, 0, W, H, WASH)
    rect(s, 0, H - Inches(0.09), W, Inches(0.09), accent)
    tf = textbox(s, Inches(1.2), Inches(2.7), Inches(11.0), Inches(2.4))
    para(tf, "Now go build.", size=44, bold=True, first=True, space_after=16)
    para(tf, "code/ — run it   ·   lab/EXERCISES.md — break it   ·   papers/ — go deeper",
         size=18, color=MUTED, space_after=8)
    para(tf, "AI End-to-End Learning Track", size=14, bold=True, color=accent)


RENDER = {"section": slide_section, "bullets": slide_bullets,
          "table": slide_table, "quote": slide_quote}


def build_pptx(deck: dict, out: Path) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    accent = hexc(deck.get("accent", "4F46E5"))
    slide_title(prs, deck, accent)
    slides = deck["slides"]
    for i, sl in enumerate(slides, start=1):
        RENDER[sl.get("type", "bullets")](prs, sl, accent, deck["module"], i, len(slides))
    slide_end(prs, deck, accent)
    prs.save(out)


def build_marp(deck: dict, out: Path) -> None:
    """Also emit a Markdown deck (Marp/Slidev compatible) so slides are diffable."""
    accent = "#" + deck.get("accent", "4F46E5").replace("#", "")
    L = ["---", "marp: true", "theme: default", "paginate: true",
         f'style: |\n  section h1 {{ color: {accent}; }}\n  section {{ font-size: 24px; }}',
         "---", "",
         f"# {deck['title']}", "", f"### {deck.get('subtitle','')}", "",
         f"**Module {deck['module']}** · AI End-to-End Learning Track", "", "---", ""]
    for sl in deck["slides"]:
        t = sl.get("type", "bullets")
        L.append(f"## {sl['title']}")
        L.append("")
        if t == "section" and sl.get("subtitle"):
            L += [f"*{sl['subtitle']}*", ""]
        if t == "quote":
            L += [f"> **{sl['quote']}**", ""]
        if t == "table":
            L.append("| " + " | ".join(sl["columns"]) + " |")
            L.append("|" + "---|" * len(sl["columns"]))
            for row in sl["rows"]:
                L.append("| " + " | ".join(str(c) for c in row) + " |")
            L.append("")
        for b in sl.get("bullets", []):
            L.append(("  - " + b[2:]) if b.startswith("- ") else ("- " + b))
        L.append("")
        if sl.get("notes"):
            L += [f"<!-- speaker note: {sl['notes']} -->", ""]
        L += ["---", ""]
    L += ["## Now go build", "", "`code/` run it · `lab/EXERCISES.md` break it · `papers/` go deeper", ""]
    out.write_text("\n".join(L), encoding="utf-8")


def main() -> int:
    targets = sys.argv[1:]
    decks = sorted(ROOT.glob("*/slides/deck.json")) + sorted(ROOT.glob("*/*/*/slides/deck.json"))
    if targets:
        decks = [d for d in decks if d.parts[-3] in targets or any(t in str(d) for t in targets)]
    if not decks:
        print("no deck.json found")
        return 1
    ok = 0
    for path in decks:
        module_dir = path.parent.parent
        try:
            deck = json.loads(path.read_text(encoding="utf-8"))
        except json.JSONDecodeError as e:
            print(f"  ✗ {module_dir.name}: bad JSON — {e}")
            continue
        name = module_dir.name
        stem = name.split("-", 1)[1] if re.match(r"^\d\d-", name) else name
        pptx_out = path.parent / f"{stem}.pptx"
        build_pptx(deck, pptx_out)
        build_marp(deck, path.parent / "slides.md")
        n = len(deck["slides"]) + 2
        print(f"  ✓ {str(module_dir.relative_to(ROOT)):<34} {n:>2} slides → {pptx_out.relative_to(ROOT)}")
        ok += 1
    print(f"\nBuilt {ok} deck(s).")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
